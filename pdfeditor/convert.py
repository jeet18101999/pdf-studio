"""Export PDF to Word (.docx), Excel (.xlsx) and text.

Word export keeps reading order, paragraphs, font family/size/bold/italic/colour, alignment guesses,
images, tables (PyMuPDF table detection) and page breaks. Excel export writes every detected table
(one sheet per page, or one combined sheet) with numbers converted to real numeric cells.
Scanned pages are exported from their OCR text layer."""
from __future__ import annotations

import io
import re
from pathlib import Path

import numpy as np
import pymupdf

from .fonts import FONT_INDEX, FLAG_BOLD, FLAG_ITALIC, color_int_to_rgb

NUM_RE = re.compile(r"^\(?[-+]?[₹$€£]?\s*\d{1,3}(?:[,\s]\d{2,3})*(?:\.\d+)?\)?%?$|^\(?[-+]?\d+(?:\.\d+)?\)?%?$")

from .cvsafe import cv2


def parse_number(text: str):
    """'1,59,500.00' -> 159500.0 ; '(1,200)' -> -1200 ; '12.5%' -> 0.125 ; otherwise None."""
    t = (text or "").strip()
    if not t or not NUM_RE.match(t):
        return None
    neg = t.startswith("(") and t.endswith(")")
    pct = t.endswith("%")
    t = re.sub(r"[()₹$€£%,\s]", "", t)
    try:
        v = float(t)
    except ValueError:
        return None
    if neg:
        v = -abs(v)
    if pct:
        v /= 100
    return int(v) if v.is_integer() and "." not in text else v


# ============================================================================ tables
def page_tables(page: pymupdf.Page):
    """Detected tables as (bbox, rows[list[list[str]]]). Vector tables first; ruled tables in scanned
    page images are detected from the image and filled with the OCR text layer."""
    out = []
    try:
        found = page.find_tables()
        for tb in found.tables:
            try:
                rows = tb.extract()
            except Exception:
                continue
            rows = [[(c or "").strip() for c in r] for r in rows]
            rows = [r for r in rows if any(c for c in r)]
            if len(rows) >= 2 and max(len(r) for r in rows) >= 2:
                out.append((pymupdf.Rect(tb.bbox), rows))
    except Exception:
        pass
    if not out:
        try:
            out = scan_tables(page)
        except Exception:
            out = []
    return out


def _is_scan(page: pymupdf.Page) -> bool:
    ur = (page.rect * page.derotation_matrix).normalize()
    return any(abs(pymupdf.Rect(i["bbox"]) & ur) > 0.6 * abs(ur) for i in page.get_image_info())


def _positions(profile, min_strength, merge):
    idx = np.nonzero(profile >= min_strength)[0]
    groups, cur = [], []
    for i in idx:
        if cur and i - cur[-1] > merge:
            groups.append(cur)
            cur = []
        cur.append(i)
    if cur:
        groups.append(cur)
    return [int(round(np.mean(g))) for g in groups]


def scan_tables(page: pymupdf.Page, dpi: int = 150):
    """Ruled tables in a scanned page image -> rows of cell text taken from the (OCR) text layer."""
    if not _is_scan(page):
        return []
    words = [w for w in page.get_text("words") if w[4].strip()]
    if len(words) < 6:
        return []
    rot = page.rotation
    page.set_rotation(0)                       # work in unrotated space, like the text layer
    try:
        pix = page.get_pixmap(dpi=dpi, alpha=False, annots=False)
        ur = page.rect
    finally:
        page.set_rotation(rot)
    k = pix.width / ur.width
    gray = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, pix.n)
    gray = cv2.cvtColor(gray, cv2.COLOR_RGB2GRAY) if pix.n >= 3 else gray[..., 0]
    _, bw = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    H, W = bw.shape
    hmask = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (max(30, W // 25), 1)))
    vmask = cv2.morphologyEx(bw, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(30, H // 45))))
    # scanned rules are broken by noise: bridge small gaps before measuring them
    hmask = cv2.morphologyEx(hmask, cv2.MORPH_CLOSE, cv2.getStructuringElement(cv2.MORPH_RECT, (max(15, W // 50), 3)))
    vmask = cv2.morphologyEx(vmask, cv2.MORPH_CLOSE, cv2.getStructuringElement(cv2.MORPH_RECT, (3, max(15, H // 70))))
    hsegs = _segments(hmask, horizontal=True, min_len=0.2 * W)     # (y, x0, x1)
    vsegs = _segments(vmask, horizontal=False, min_len=0.08 * H)   # (x, y0, y1)
    if len(vsegs) < 2:
        return []

    # a table = column dividers whose vertical ranges overlap (linked transitively); page-frame lines are
    # never used to link sections, only as the table's outer edges
    frame = [s for s in vsegs if s[2] - s[1] > 0.6 * H]
    inner = [s for s in vsegs if s[2] - s[1] <= 0.6 * H]
    parent = list(range(len(inner)))

    def find(i):
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i
    for i, a in enumerate(inner):
        for j in range(i + 1, len(inner)):
            b = inner[j]
            overlap = min(a[2], b[2]) - max(a[1], b[1])
            if overlap > 0.6 * min(a[2] - a[1], b[2] - b[1]):
                parent[find(i)] = find(j)
    groups: dict[int, list[tuple]] = {}
    for i, s in enumerate(inner):
        groups.setdefault(find(i), []).append(s)

    tables = []
    for g in sorted(groups.values(), key=len, reverse=True):
        if len(g) < 2:
            continue
        y0 = int(min(s[1] for s in g))
        y1 = int(max(s[2] for s in g))
        span = y1 - y0
        if span < 0.05 * H:
            continue
        gx0, gx1 = min(s[0] for s in g), max(s[0] for s in g)
        covering = [s[0] for s in frame + inner if s[1] <= y0 + 0.15 * span and s[2] >= y1 - 0.15 * span]
        left_c = max([x for x in covering if x < gx0], default=None)
        right_c = min([x for x in covering if x > gx1], default=None)
        # horizontal rules of this section tell where it really starts/ends (e.g. a grid in the right half)
        band_rules = [(y, a, b) for y, a, b in hsegs if y0 - 0.15 * span <= y <= y1 + 8 and a < gx1 and b > gx0]
        if band_rules:
            starts = sorted(a for _, a, _ in band_rules)
            ends = sorted(b for _, _, b in band_rules)
            rule_left, rule_right = starts[len(starts) // 2], ends[len(ends) // 2]
            if left_c is None or rule_left > left_c + 12:
                left_c = min(rule_left, gx0 - 4)
            if right_c is None or rule_right < right_c - 12:
                right_c = max(rule_right, gx1 + 4)
        xs = [s[0] for s in g]
        xs += [x for x in covering if (left_c or 0) <= x <= (right_c or W)]
        if left_c is not None:
            xs.append(left_c)
        if right_c is not None:
            xs.append(right_c)
        xs = _merge_close(xs, 6)
        if len(xs) < 3 or xs[-1] - xs[0] < 0.25 * W:
            continue
        rules_y = sorted(y for y, a, b in band_rules) if band_rules else []
        top = min([y0] + [y for y in rules_y if y < y0])
        bbox = pymupdf.Rect(xs[0] / k, top / k, xs[-1] / k, y1 / k)
        if any(abs(bbox & t[0]) > 0.5 * abs(bbox) for t in tables):
            continue
        rows = _fill_table(words, bbox, [x / k for x in xs], [y / k for y in rules_y])
        if len(rows) >= 2 and max(len(r) for r in rows) >= 2:
            tables.append((bbox, rows))
    tables.sort(key=lambda t: (t[0].y0, t[0].x0))
    return tables


def _segments(mask, horizontal: bool, min_len: float):
    n, _, stats, _ = cv2.connectedComponentsWithStats(mask)
    out = []
    for i in range(1, n):
        x, y, w, h, _ = stats[i]
        if horizontal and w >= min_len and h <= 12:
            out.append((y + h // 2, x, x + w))
        elif not horizontal and h >= min_len and w <= 12:
            out.append((x + w // 2, y, y + h))
    return out


def _merge_close(values, dist):
    values = sorted(values)
    out = []
    for v in values:
        if out and v - out[-1] <= dist:
            continue
        out.append(v)
    return out


def _fill_table(words, bbox: pymupdf.Rect, col_edges, row_rules):
    inside = [w for w in words if bbox.contains(pymupdf.Point((w[0] + w[2]) / 2, (w[1] + w[3]) / 2))]
    if not inside:
        return []
    ncols = max(1, len(col_edges) - 1)

    def col_of(xc):
        for c in range(ncols):
            if col_edges[c] - 1 <= xc <= col_edges[c + 1] + 1:
                return c
        return 0 if xc < col_edges[0] else ncols - 1

    # text lines by baseline clustering
    inside.sort(key=lambda w: ((w[1] + w[3]) / 2, w[0]))
    lines, cur, cur_y = [], [], None
    for w in inside:
        yc = (w[1] + w[3]) / 2
        h = w[3] - w[1]
        if cur and abs(yc - cur_y) > max(3.0, 0.5 * h):
            lines.append(cur)
            cur = []
        cur.append(w)
        cur_y = yc if len(cur) == 1 else (cur_y * (len(cur) - 1) + yc) / len(cur)
    if cur:
        lines.append(cur)

    rows: list[list[str]] = []
    row_band: list[int] = []          # which ruled band each row belongs to
    for ln in lines:
        yc = sum((w[1] + w[3]) / 2 for w in ln) / len(ln)
        band = sum(1 for r in row_rules if r < yc)
        cells = [[] for _ in range(ncols)]
        for w in sorted(ln, key=lambda w: w[0]):
            cells[col_of((w[0] + w[2]) / 2)].append(w[4])
        texts = [" ".join(c) for c in cells]
        filled = [i for i, t in enumerate(texts) if t]
        if rows and row_band[-1] == band and len(filled) == 1:
            prev = rows[-1]
            c = filled[0]
            prev_filled = [i for i, t in enumerate(prev) if t]
            # continuation line of a multi-line cell (e.g. item description): merge into the row above
            if len(prev_filled) >= 2 or (prev_filled == [c] and band_rows_single(row_band, band)):
                prev[c] = (prev[c] + "\n" + texts[c]).strip()
                continue
        rows.append(texts)
        row_band.append(band)
    # drop fully empty columns
    keep = [c for c in range(ncols) if any(r[c] for r in rows)]
    return [[r[c] for c in keep] for r in rows]


def band_rows_single(row_band, band) -> bool:
    """Inside a ruled cell band with a single logical row, stacked lines belong together."""
    return row_band.count(band) == 1


# ============================================================================ Excel
def to_xlsx(doc: pymupdf.Document, path: str, one_sheet: bool = False, pages=None) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_fill = PatternFill("solid", fgColor="E8EEFC")
    count = 0
    combined = wb.create_sheet("Tables") if one_sheet else None
    crow = 1
    for pno in (pages if pages is not None else range(doc.page_count)):
        tables = page_tables(doc[pno])
        if not tables:
            continue
        ws = combined if one_sheet else wb.create_sheet(f"Page {pno + 1}")
        row = crow if one_sheet else 1
        for ti, (bbox, rows) in enumerate(tables):
            count += 1
            ws.cell(row=row, column=1, value=f"Page {pno + 1} – Table {ti + 1}").font = Font(bold=True, color="1F3A93")
            row += 1
            for ri, r in enumerate(rows):
                for ci, text in enumerate(r):
                    num = parse_number(text)
                    cell = ws.cell(row=row, column=ci + 1, value=num if num is not None else text)
                    cell.border = border
                    cell.alignment = Alignment(wrap_text=True, vertical="top",
                                               horizontal="right" if num is not None else "left")
                    if num is not None and isinstance(num, float):
                        cell.number_format = "#,##0.00"
                    if ri == 0:
                        cell.font = Font(bold=True)
                        cell.fill = header_fill
                row += 1
            row += 1
        # column widths
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None and c.row > 0), default=8)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(60, max(8, width * 0.9 + 2))
        if one_sheet:
            crow = row
    if not wb.sheetnames:
        ws = wb.create_sheet("Text")
        for pno in range(doc.page_count):
            for i, line in enumerate(doc[pno].get_text().splitlines()):
                ws.append([line])
    wb.save(path)
    return count


# ============================================================================ Word
def _docx_color(rgb):
    from docx.shared import RGBColor
    return RGBColor(int(rgb[0] * 255), int(rgb[1] * 255), int(rgb[2] * 255))


def _family(font_name: str, flags: int) -> str:
    try:
        entry, _, _ = FONT_INDEX.match_pdf_font(font_name, flags)
        return entry.family if entry else "Arial"
    except Exception:
        return "Arial"


def to_docx(doc: pymupdf.Document, path: str, pages=None, progress=None) -> str:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt

    FONT_INDEX.ensure()
    out = Document()
    pages = list(pages if pages is not None else range(doc.page_count))
    for n, pno in enumerate(pages):
        page = doc[pno]
        rect = page.rect
        section = out.sections[-1] if n == 0 else out.add_section()
        section.page_width, section.page_height = Pt(rect.width), Pt(rect.height)
        section.orientation = WD_ORIENT.LANDSCAPE if rect.width > rect.height else WD_ORIENT.PORTRAIT
        margin = Pt(36)
        section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = margin
        usable = rect.width - 72

        tables = page_tables(page)
        items = []           # (y, kind, payload)
        for bbox, rows in tables:
            items.append((bbox.y0, "table", rows))
        d = page.get_text("dict", flags=pymupdf.TEXTFLAGS_DICT & ~pymupdf.TEXT_PRESERVE_IMAGES | pymupdf.TEXT_PRESERVE_IMAGES)
        for b in d["blocks"]:
            bb = pymupdf.Rect(b["bbox"])
            if any(abs(bb & t[0]) > 0.5 * abs(bb) for t in tables if abs(bb) > 0):
                continue            # text inside a table is exported with the table
            if b["type"] == 1:
                items.append((bb.y0, "image", (b, bb)))
            elif b["type"] == 0:
                items.append((bb.y0, "text", (b, bb)))
        items.sort(key=lambda it: it[0])

        for _, kind, payload in items:
            if kind == "table":
                rows = payload
                ncols = max(len(r) for r in rows)
                t = out.add_table(rows=len(rows), cols=ncols)
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, r in enumerate(rows):
                    for ci in range(ncols):
                        text = r[ci] if ci < len(r) else ""
                        cell = t.cell(ri, ci)
                        cell.text = text
                        p = cell.paragraphs[0]
                        if parse_number(text) is not None:
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        for run in p.runs:
                            run.font.size = Pt(9)
                            run.font.bold = ri == 0
                out.add_paragraph()
            elif kind == "image":
                b, bb = payload
                try:
                    pic = b.get("image")
                    if pic:
                        width = min(usable, bb.width)
                        out.add_picture(io.BytesIO(pic), width=Pt(width))
                        out.paragraphs[-1].alignment = _align_for(bb, rect)
                except Exception:
                    pass
            else:
                b, bb = payload
                para = out.add_paragraph()
                para.paragraph_format.space_after = Pt(4)
                indent = max(0.0, bb.x0 - 36)
                if indent > 4 and indent < usable * 0.6:
                    para.paragraph_format.left_indent = Pt(indent)
                para.alignment = _align_for(bb, rect)
                lines = [l for l in b["lines"] if any(s["text"].strip() for s in l["spans"])]
                for li, line in enumerate(lines):
                    for s in line["spans"]:
                        if not s["text"]:
                            continue
                        run = para.add_run(s["text"])
                        run.font.name = _family(s["font"], s["flags"])
                        run.font.size = Pt(max(1.0, round(s["size"] * 2) / 2))
                        run.font.bold = bool(s["flags"] & FLAG_BOLD) or "bold" in s["font"].lower()
                        run.font.italic = bool(s["flags"] & FLAG_ITALIC) or "italic" in s["font"].lower()
                        if s.get("alpha", 255) == 0:
                            pass               # invisible OCR text: keep the words, default colour
                        elif s["color"]:
                            run.font.color.rgb = _docx_color(color_int_to_rgb(s["color"]))
                    if li < len(lines) - 1:
                        # join wrapped lines of a paragraph with a space; keep hard breaks for short lines
                        cur_w = pymupdf.Rect(line["bbox"]).width
                        if cur_w < bb.width * 0.75:
                            para.add_run().add_break(WD_BREAK.LINE)
                        else:
                            para.add_run(" ")
        if progress:
            progress(n + 1, len(pages))
    out.save(path)
    return path


def _align_for(bb: pymupdf.Rect, page_rect: pymupdf.Rect):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    left = bb.x0 - page_rect.x0
    right = page_rect.x1 - bb.x1
    if abs(left - right) < 20 and left > 80:
        return WD_ALIGN_PARAGRAPH.CENTER
    if right < 60 and left > page_rect.width * 0.45:
        return WD_ALIGN_PARAGRAPH.RIGHT
    return WD_ALIGN_PARAGRAPH.LEFT


# ============================================================================ text
def to_text(doc: pymupdf.Document, path: str, layout: bool = True) -> str:
    chunks = []
    for page in doc:
        if layout:
            try:
                chunks.append(page.get_text("text", sort=True))
                continue
            except Exception:
                pass
        chunks.append(page.get_text())
    Path(path).write_text("\n\f".join(chunks), "utf-8")
    return path
